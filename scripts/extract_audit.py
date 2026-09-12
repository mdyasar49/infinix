import docx

doc_path = r"d:\infonix\InfogenX Audiot report.docx"
doc = docx.Document(doc_path)

with open(r"d:\infonix\scripts\extracted_audit_report.txt", "w", encoding="utf-8") as out:
    out.write("=== PARAGRAPHS ===\n")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            out.write(f"P{i+1}: {p.text.strip()}\n")
    
    out.write("\n=== TABLES ===\n")
    for t_idx, t in enumerate(doc.tables):
        out.write(f"\n--- TABLE {t_idx+1} ---\n")
        for r_idx, row in enumerate(t.rows):
            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            out.write(f"R{r_idx+1}: " + " | ".join(cells) + "\n")

print("Extracted successfully.")
